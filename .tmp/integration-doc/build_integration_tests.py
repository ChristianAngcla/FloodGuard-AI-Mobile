from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.table import Table
from docx.shared import Pt
from docx.oxml.ns import qn

SRC = Path(r"C:\Users\inugami\Downloads\CCIT-RESEARCH-TEST Integration Test Document.docx")
OUT = Path(r"C:\Users\inugami\Desktop\FloodGuard-AI-Mobile\output\FloodGuard_AI_Integration_Test_Document_All_Modules.docx")

cases = [
 ("IT-001", "App Startup -> Firebase Core -> Notification Service -> SharedPreferences -> Home Map",
  "Verify that startup services initialize safely, saved preferences load, and the mobile app opens the map without blocking on optional services.",
  "Install a test build; prepare clean and previously used app states; internet may be enabled or disabled.",
  [("Cold-launch after fresh install and respond to the notification prompt.", "Fresh app data; notifications allowed", "Firebase and notifications initialize; default light/English settings load; Home Map opens."),
   ("Cold-launch after fresh install and deny notifications.", "Fresh app data; notifications denied", "Home Map still opens; denial does not crash or block the app."),
   ("Relaunch after enabling dark mode and Taglish.", "is_dark_mode=true; is_taglish=true", "Saved theme and language are restored on launch."),
   ("Launch with mobile data/Wi-Fi disabled.", "Offline device", "App opens with local map assets and handles unavailable remote data gracefully."),
   ("Background the app, terminate it, then relaunch.", "Existing preferences and cached session", "Startup completes once and returns to a stable Home Map state.")]),
 ("IT-002", "Sign-Up Screen -> Firebase Phone OTP -> Auth API -> User Database -> Local Preferences",
  "Verify that verified registration details are sent to the FloodGuard API, stored as a citizen account, and made available to the mobile session.",
  "Use an unregistered email and Philippine mobile number; Firebase phone authentication and FloodGuard API are reachable.",
  [("Complete all registration steps, request OTP, enter the valid code, and submit.", "Unique valid resident data", "OTP is verified, the API creates one account, and the app proceeds to the authenticated flow."),
   ("Request OTP and enter an invalid code.", "Incorrect six-digit OTP", "Firebase rejects the code and no user record is created."),
   ("Submit registration using an email already registered.", "Existing email; valid phone OTP", "API rejects the duplicate account and the UI shows a clear message."),
   ("Submit a valid account using normalized address and phone values.", "09xxxxxxxxx; Marikina address", "API receives complete profile fields and stores normalized resident data."),
   ("Interrupt the network after OTP verification, restore it, and retry once.", "Unique account; temporary network loss", "The UI reports the connectivity issue and retry does not create duplicate accounts.")]),
 ("IT-003", "Login Screen -> Auth API -> User Database -> JWT -> SharedPreferences -> Home/Profile",
  "Verify authentication, token and profile persistence, and routing into the authenticated mobile experience.",
  "A verified citizen account exists in the user database; FloodGuard API is reachable.",
  [("Log in with valid email and password.", "Registered credentials", "API returns success, JWT and user data are stored, and Home Map opens."),
   ("Open Profile immediately after successful login.", "Stored JWT and user_data", "Profile displays the same resident identity returned by the API."),
   ("Log in with a valid email and incorrect password.", "Registered email; wrong password", "API denies access; no authenticated session or token is stored."),
   ("Submit credentials while the API is unavailable.", "Valid credentials; offline/API down", "The UI shows failure without crashing or creating a false session."),
   ("Terminate and relaunch after a successful login.", "Persisted is_logged_in, token, and user_data", "The app restores authenticated profile access from the saved session.")]),
 ("IT-004", "Forgot Password -> Account Lookup API -> Firebase Phone Auth -> Password API -> Login",
  "Verify the complete password-recovery flow from account lookup and phone verification through password update and subsequent login.",
  "A registered account has an accessible phone number; Firebase and FloodGuard authentication endpoints are available.",
  [("Look up a registered email, verify the displayed masked phone, and request OTP.", "Registered email", "The correct account is found and OTP is sent to the linked phone."),
   ("Enter a valid OTP and set a compliant new password.", "Valid OTP; new password", "Password update succeeds and confirmation is displayed."),
   ("Attempt recovery with an unknown email.", "Unregistered email", "No account is disclosed; the user receives a clear not-found response."),
   ("Enter an invalid or expired OTP.", "Invalid/expired OTP", "Verification fails and the password remains unchanged."),
   ("Log in using the new password, then try the old password.", "Updated and previous passwords", "New password authenticates; old password is rejected.")]),
 ("IT-005", "Profile Screen -> Local Session UID -> Profile API -> User Database -> FCM Topic Subscription",
  "Verify profile retrieval and editing, database persistence, local cache synchronization, and barangay alert subscription updates.",
  "User is logged in; profile record exists; API and FCM are available; notification permission is allowed.",
  [("Open Profile from an authenticated session.", "Stored user_data and UID", "Profile fields load consistently from the saved/API-backed identity."),
   ("Edit valid contact and address fields and save.", "Valid profile updates", "API/database save succeeds and refreshed Profile shows the updates."),
   ("Change the selected barangay and save.", "Different Marikina barangay", "Profile persists and the device subscribes to the new barangay topic."),
   ("Enter an invalid phone value and attempt to save.", "Malformed/short phone", "Client validation blocks the API update and explains the error."),
   ("Restart the app and reopen Profile after a successful update.", "Previously saved profile", "Updated values remain synchronized across local cache and backend data.")]),
 ("IT-006", "Logout -> FCM Topic Unsubscribe -> SharedPreferences Session Clear -> Login/Profile Access",
  "Verify logout removes authentication state, ends barangay notification subscription, and prevents authenticated profile access.",
  "User is logged in and subscribed to a barangay notification topic.",
  [("Select Logout from Profile and confirm.", "Active authenticated session", "FCM topic is unsubscribed, token/user data are cleared, and Login opens."),
   ("Press Back after logout.", "Logged-out state", "Authenticated Profile cannot be restored from navigation history."),
   ("Terminate and relaunch after logout.", "Cleared session", "No stale authenticated profile is restored."),
   ("Log out while the notification service is unavailable.", "Active session; FCM failure", "Local authentication data is still cleared and the user exits safely."),
   ("Log in again after logout.", "Valid credentials", "A new valid session is created without stale profile or subscription data.")]),
 ("IT-007", "Home Map -> Local GeoJSON -> Flood API /status -> Sensor Mapping -> Risk Polygons and Legend",
  "Verify local barangay boundaries combine with live flood telemetry and station thresholds to render the correct risk state on the map.",
  "GeoJSON assets are bundled; FloodGuard API status endpoint contains mapped sensor data.",
  [("Open Home Map with the API online and wait for initial loading.", "Live /api/status response", "All supported barangay polygons load and show risk colors matching mapped sensor status."),
   ("Tap a barangay polygon.", "Known mapped barangay", "The selected barangay opens and displays the same status, rainfall, and water-level data."),
   ("Pull/trigger refresh after telemetry changes.", "Updated API timestamp/status", "Map colors and displayed measurements update without duplicate polygons."),
   ("Load the map when one sensor value is missing.", "Partial live_sensors payload", "Affected areas show unavailable/fallback state; other barangays remain accurate."),
   ("Load while the API is offline.", "GeoJSON available; API unavailable", "Boundaries remain usable and remote data failure is communicated without a crash.")]),
 ("IT-008", "Location Permission -> Geolocator Stream -> Map Camera -> Weather Coordinates",
  "Verify device location permission and GPS updates feed the map position and coordinate-dependent weather workflow.",
  "Run on an emulator/device with controllable GPS and permission settings.",
  [("Allow location access and tap Center on Me.", "GPS inside Marikina", "Current-location marker appears and camera centers on the reported coordinates."),
   ("Move the emulated location by more than five meters.", "Second valid coordinate", "Position stream updates the marker without freezing map navigation."),
   ("Deny location permission.", "Permission denied", "Map remains usable and the app gives a non-crashing permission response."),
   ("Disable device location services and request centering.", "GPS service disabled", "The app does not claim a location and handles the disabled service safely."),
   ("Restore permission and GPS, then retry.", "GPS enabled; permission granted", "Location marker, camera, and location-based weather request recover correctly.")]),
 ("IT-009", "Weather Card -> Open-Meteo API -> Weather Parser/Cache -> Dashboard UI",
  "Verify coordinates are converted into current weather and five-day forecast data, cached, translated, and displayed correctly.",
  "Internet access is available; Open-Meteo endpoint is reachable; valid coordinates are supplied.",
  [("Open the weather card for Marikina coordinates.", "14.6507, 121.1029", "Current temperature, humidity, precipitation, wind, heat index, rain chance, and forecast render."),
   ("Refresh/open again within ten minutes using the same coordinates.", "Same rounded coordinate key", "Cached weather is reused and displayed values remain consistent."),
   ("Request weather for a second coordinate.", "Different valid coordinate", "A new request is made and the UI displays the second location's response."),
   ("Simulate an API timeout or non-200 response.", "Unavailable Open-Meteo", "Weather UI shows an unavailable/retry state without affecting flood map data."),
   ("Switch app language after weather loads.", "English then Taglish", "Weather code labels and surrounding UI update consistently without corrupting values.")]),
 ("IT-010", "Barangay Details/Dashboard -> Flood Forecast API -> Sensor Mapping -> Forecast Cards/Charts",
  "Verify selected barangay and mapped river station retrieve and present consistent current and daily forecast information.",
  "FloodGuard API provides status and /forecasts/daily data; barangay-to-sensor mapping is configured.",
  [("Open details for a mapped barangay from the map.", "Barangay with live sensor", "Correct river station, current level, threshold, and status are shown."),
   ("Change the selected barangay in the dashboard.", "Second mapped barangay", "All cards/charts update to the second barangay and its mapped station."),
   ("Load daily forecasts for the selected barangay.", "Valid forecast payload", "Forecast date, predicted level, mode, and status fields display consistently."),
   ("Use a payload marked persistence fallback.", "calculationMode=persistence_fallback", "UI labels the fallback mode and does not present it as primary-model output."),
   ("Use unavailable or incomplete forecast data.", "Missing predictedWaterLevel", "UI displays forecast unavailable without fabricated measurements or a crash.")]),
 ("IT-011", "Flood Report Sheet -> Authenticated UID -> Image Picker/GPS -> Flood API -> Report Database",
  "Verify a multi-step citizen report combines identity, barangay, observations, optional evidence, and location into one backend submission.",
  "User is logged in; report endpoint is available; prepare a test image and controllable GPS permission.",
  [("Complete all report steps with valid observations and submit.", "UID, barangay, depth/category, notes", "API accepts one report and the UI confirms successful submission."),
   ("Attach a valid image and submit with location allowed.", "JPEG/PNG evidence plus GPS", "Evidence/location are associated with the submitted report without losing form data."),
   ("Submit without optional image evidence.", "Valid required fields; no image", "Report is accepted if all required fields are present."),
   ("Attempt to continue with required fields missing.", "Missing barangay or report detail", "Client blocks submission and identifies the missing input."),
   ("Lose connectivity during submission, restore it, and retry once.", "Valid report; temporary offline state", "Failure is explained and retry creates no unintended duplicate report.")]),
 ("IT-012", "Backend Alert -> FCM Barangay Topic -> Local Notification -> SharedPreferences -> Alerts Screen",
  "Verify targeted flood alerts reach subscribed devices, persist once, display locally, and appear in the in-app alert history.",
  "Device has notification permission and an active barangay subscription; a test FCM sender/backend is available.",
  [("Send a test alert to the user's subscribed barangay while app is foregrounded.", "Unique messageId; title/body/data", "High-priority local notification appears and one unread alert is persisted."),
   ("Open Alerts after receiving the message.", "Persisted app_alerts entry", "Alert title, body, timestamp, and unread state are displayed correctly."),
   ("Send the same messageId twice.", "Duplicate FCM message", "Only one history entry is stored."),
   ("Send an alert while the app is backgrounded, then open it.", "Unique background message", "Background handler persists the alert and opening the app exposes it in Alerts."),
   ("Send an alert for a different barangay topic.", "Non-subscribed topic", "The device does not receive or store the unrelated alert.")]),
 ("IT-013", "App Drawer/Navigation -> Map, Home Dashboard, Alerts, Profile, Login/Sign-Up",
  "Verify navigation connects all principal screens while preserving shared session, map, theme, and language state.",
  "App is installed; conduct cycles both logged out and logged in.",
  [("Navigate Map -> Dashboard -> Alerts -> Map using the drawer.", "Logged-out or guest state", "Each destination opens once and Back/drawer behavior remains stable."),
   ("Open Profile while logged out.", "No authenticated session", "App routes to Login or restricted profile flow without exposing resident data."),
   ("Log in, then navigate to Profile and back to Map.", "Valid authenticated session", "Authenticated profile opens and map state remains usable on return."),
   ("Open Login, Sign-Up, and return using screen controls.", "Logged-out state", "Authentication screens transition without stacking duplicate routes."),
   ("Navigate repeatedly in dark/Taglish mode.", "Saved dark mode and Taglish", "Theme/language remain consistent across every destination.")]),
 ("IT-014", "Theme/Language Controls -> SharedPreferences -> App Theme/Translations -> All Screens",
  "Verify appearance and language selections propagate across modules and persist across app restarts.",
  "All principal screens are reachable; local preferences storage is available.",
  [("Switch from light to dark mode on the map.", "is_dark_mode=true", "Map chrome, drawer, cards, dialogs, and text adopt dark styling without unreadable content."),
   ("Navigate through Dashboard, Alerts, Profile, Login, and Sign-Up in dark mode.", "Dark mode active", "All screens receive the same theme state."),
   ("Switch from English to Taglish.", "is_taglish=true", "Supported labels/messages update across visible modules."),
   ("Restart the app after changing both settings.", "Persisted theme/language values", "The selected appearance and language restore at startup."),
   ("Toggle both settings repeatedly while navigating.", "Alternating preference values", "No stale mixed-theme/mixed-language screen or crash occurs.")]),
]

def set_text(cell, text, bold=False, size=7.5):
    cell.text = str(text)
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.name = "Arial Narrow"
            r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial Narrow")
            r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial Narrow")
            r.font.size = Pt(size)
            r.bold = bold

def fill_form(tbl_el, case):
    t = Table(tbl_el, doc)
    it_id, points, desc, pre, cycles = case
    set_text(t.cell(0, 0), "INTEGRATION TEST DOCUMENT", True, 10)
    set_text(t.cell(0, 2), points, True, 8)
    set_text(t.cell(1, 1), it_id, True, 8)
    set_text(t.cell(1, 2), desc, False, 7.5)
    set_text(t.cell(2, 1), "____________________________", False, 7.5)
    set_text(t.cell(2, 3), "Mobile Application (Android)", False, 7.5)
    set_text(t.cell(3, 1), pre, False, 7.5)
    headers = ["Test Cycle", "Steps/Actions", "Steps/Actions", "Data", "Expected Results", "Actual Results", "Pass/Fail"]
    for col, header in enumerate(headers):
        set_text(t.cell(4, col), header, True, 7.5)
    for idx, (action, data, expected) in enumerate(cycles):
        row = 5 + idx
        set_text(t.cell(row, 0), str(idx + 1), True, 7.5)
        set_text(t.cell(row, 1), action, False, 7.2)
        set_text(t.rows[row].cells[2], data, False, 7.2)
        set_text(t.rows[row].cells[3], expected, False, 7.2)
        set_text(t.rows[row].cells[4], "", False, 7.2)
        set_text(t.rows[row].cells[5], "", False, 7.2)
    set_text(t.rows[10].cells[0], "Comments/Remarks: To be completed by the assigned tester during execution.", False, 7.5)

doc = Document(SRC)
body = doc._element.body
source_test = deepcopy(doc.tables[0]._tbl)
source_sig = deepcopy(doc.tables[1]._tbl)
sect_pr = body.sectPr
for child in list(body):
    if child is not sect_pr:
        body.remove(child)

for idx, case in enumerate(cases):
    test_el = deepcopy(source_test)
    body.insert(len(body)-1, test_el)
    fill_form(test_el, case)
    sig_el = deepcopy(source_sig)
    body.insert(len(body)-1, sig_el)
    sig = Table(sig_el, doc)
    set_text(sig.cell(0,0), "Prepared By", True, 8)
    set_text(sig.cell(0,1), "Administered/Performed By", True, 8)
    set_text(sig.cell(1,0), "____________________________\nSignature over Printed Name", False, 7.5)
    set_text(sig.cell(1,1), "____________________________\nSignature over Printed Name", False, 7.5)
    if idx != len(cases)-1:
        p = doc.add_page_break()
        body.remove(p._p)
        body.insert(len(body)-1, p._p)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = "FloodGuard AI Integration Test Document - All Modules"
doc.core_properties.subject = "Alpha-stage integration test plan and execution forms"
doc.core_properties.author = "FloodGuard AI Research Team"
doc.save(OUT)
print(OUT)
